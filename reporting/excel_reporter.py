"""
Excel report generation module
"""

import pandas as pd
import numpy as np
import logging
from typing import Dict, List, Optional, Any
from datetime import datetime
import os

# Excel libraries
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
import xlsxwriter

from reporting.visualizations import ChartGenerator
from utils.helpers import format_currency, format_percentage, format_number

class ExcelReporter:
    """Class for generating comprehensive Excel reports"""
    
    def __init__(self, config):
        self.config = config
        self.logger = logging.getLogger(__name__)
        self.chart_generator = ChartGenerator(config)
        
    def generate_report(self, processed_data: Dict[str, pd.DataFrame],
                       exposure_results: Dict[str, Any],
                       hedge_ratios: Dict[str, Any],
                       backtest_results: Dict[str, Any],
                       ml_results: Optional[Dict[str, Any]] = None) -> str:
        """Generate comprehensive Excel report"""
        
        report_path = self.config.get_report_path()
        self.logger.info(f"Generating Excel report: {report_path}")
        
        try:
            # Create Excel writer
            with pd.ExcelWriter(report_path, engine='xlsxwriter') as writer:
                
                # Get workbook and add formats
                workbook = writer.book
                self._add_formats(workbook)
                
                # Generate worksheets
                self._create_executive_summary(writer, exposure_results, hedge_ratios, backtest_results, ml_results)
                self._create_data_summary(writer, processed_data)
                self._create_exposure_analysis(writer, exposure_results)
                self._create_hedge_ratios_analysis(writer, hedge_ratios)
                self._create_backtesting_results(writer, backtest_results)
                self._create_risk_analysis(writer, backtest_results)
                self._create_performance_comparison(writer, backtest_results)
                self._create_scenario_analysis(writer, backtest_results)
                self._create_detailed_data(writer, processed_data)
                
                # Add ML results if available
                if ml_results and ml_results.get('status') == 'completed':
                    self._create_ml_analysis(writer, ml_results)
                    self._create_ml_predictions(writer, ml_results)
                    self._create_ml_model_comparison(writer, ml_results)
                
                # Add charts if enabled
                if self.config.INCLUDE_CHARTS:
                    self._add_charts_to_report(writer, processed_data, backtest_results, ml_results)
            
            self.logger.info(f"Excel report generated successfully: {report_path}")
            return report_path
            
        except Exception as e:
            self.logger.error(f"Error generating Excel report: {e}")
            raise
    
    def _add_formats(self, workbook):
        """Add custom formats to workbook"""
        self.formats = {
            'title': workbook.add_format({
                'font_size': 16,
                'bold': True,
                'align': 'center',
                'valign': 'vcenter',
                'bg_color': '#4472C4',
                'font_color': 'white'
            }),
            'header': workbook.add_format({
                'font_size': 12,
                'bold': True,
                'align': 'center',
                'valign': 'vcenter',
                'bg_color': '#D9E2F3',
                'border': 1
            }),
            'subheader': workbook.add_format({
                'font_size': 11,
                'bold': True,
                'bg_color': '#E7E6E6',
                'border': 1
            }),
            'currency': workbook.add_format({
                'num_format': '$#,##0.00',
                'border': 1
            }),
            'percentage': workbook.add_format({
                'num_format': '0.00%',
                'border': 1
            }),
            'number': workbook.add_format({
                'num_format': '#,##0.00',
                'border': 1
            }),
            'date': workbook.add_format({
                'num_format': 'yyyy-mm-dd',
                'border': 1
            }),
            'positive': workbook.add_format({
                'num_format': '0.00%',
                'font_color': 'green',
                'border': 1
            }),
            'negative': workbook.add_format({
                'num_format': '0.00%',
                'font_color': 'red',
                'border': 1
            }),
            'cell_border': workbook.add_format({'border': 1})
        }
    
    def _create_executive_summary(self, writer, exposure_results, hedge_ratios, backtest_results, ml_results=None):
        """Create executive summary worksheet"""
        worksheet = writer.book.add_worksheet('Executive Summary')
        
        # Title
        worksheet.merge_range('A1:H1', 'FBX Hedging Strategy - Executive Summary', self.formats['title'])
        worksheet.set_row(0, 25)
        
        # Report metadata
        row = 2
        worksheet.write(row, 0, 'Report Date:', self.formats['subheader'])
        worksheet.write(row, 1, datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        row += 1
        
        worksheet.write(row, 0, 'Analysis Period:', self.formats['subheader'])
        worksheet.write(row, 1, f"{self.config.START_DATE.strftime('%Y-%m-%d')} to {self.config.END_DATE.strftime('%Y-%m-%d')}")
        row += 2
        
        # Key findings
        worksheet.write(row, 0, 'KEY FINDINGS', self.formats['header'])
        row += 1
        
        # Exposure analysis summary
        total_exposure = exposure_results.get('total_exposure', {})
        revenue_sensitivity = total_exposure.get('revenue_sensitivity', 0)
        dollar_exposure = total_exposure.get('dollar_exposure', 0)
        
        worksheet.write(row, 0, '1. Revenue Exposure to FBX:', self.formats['subheader'])
        row += 1
        worksheet.write(row, 1, f"Revenue Sensitivity: {format_percentage(revenue_sensitivity)}")
        row += 1
        worksheet.write(row, 1, f"Dollar Exposure: {format_currency(dollar_exposure)}")
        row += 2
        
        # Hedge strategy performance
        worksheet.write(row, 0, '2. Hedge Strategy Performance:', self.formats['subheader'])
        row += 1
        
        # Get best performing strategy
        best_strategy = self._get_best_strategy(backtest_results)
        if best_strategy:
            strategy_name, strategy_results = best_strategy
            total_return = strategy_results.get('total_return', 0)
            hedge_effectiveness = backtest_results.get('hedge_effectiveness', {}).get(strategy_name, {}).get('hedge_effectiveness', 0)
            
            worksheet.write(row, 1, f"Best Strategy: {strategy_name.replace('_', ' ').title()}")
            row += 1
            worksheet.write(row, 1, f"Total Return: {format_percentage(total_return)}")
            row += 1
            worksheet.write(row, 1, f"Hedge Effectiveness: {format_percentage(hedge_effectiveness)}")
            row += 2
        
        # Risk metrics
        worksheet.write(row, 0, '3. Risk Metrics:', self.formats['subheader'])
        row += 1
        
        unhedged_results = backtest_results.get('unhedged', {})
        unhedged_vol = unhedged_results.get('metrics', {}).get('volatility', 0)
        unhedged_mdd = unhedged_results.get('metrics', {}).get('max_drawdown', 0)
        
        worksheet.write(row, 1, f"Unhedged Volatility: {format_percentage(unhedged_vol)}")
        row += 1
        worksheet.write(row, 1, f"Unhedged Max Drawdown: {format_percentage(unhedged_mdd)}")
        row += 1
        
        if best_strategy:
            hedged_vol = strategy_results.get('metrics', {}).get('volatility', 0)
            hedged_mdd = strategy_results.get('metrics', {}).get('max_drawdown', 0)
            vol_reduction = (unhedged_vol - hedged_vol) / unhedged_vol if unhedged_vol != 0 else 0
            
            worksheet.write(row, 1, f"Hedged Volatility: {format_percentage(hedged_vol)}")
            row += 1
            worksheet.write(row, 1, f"Volatility Reduction: {format_percentage(vol_reduction)}")
            row += 2
        
        # Recommendations
        worksheet.write(row, 0, '4. Recommendations:', self.formats['subheader'])
        row += 1
        
        recommendations = self._generate_recommendations(exposure_results, hedge_ratios, backtest_results)
        for i, rec in enumerate(recommendations[:5], 1):
            worksheet.write(row, 1, f"{i}. {rec}")
            row += 1
        
        # Format columns
        worksheet.set_column('A:A', 25)
        worksheet.set_column('B:H', 15)
    
    def _create_data_summary(self, writer, processed_data):
        """Create data summary worksheet"""
        worksheet = writer.book.add_worksheet('Data Summary')
        
        # Title
        worksheet.merge_range('A1:F1', 'Data Summary', self.formats['title'])
        
        row = 2
        
        # Data overview
        worksheet.write(row, 0, 'Dataset', self.formats['header'])
        worksheet.write(row, 1, 'Observations', self.formats['header'])
        worksheet.write(row, 2, 'Start Date', self.formats['header'])
        worksheet.write(row, 3, 'End Date', self.formats['header'])
        worksheet.write(row, 4, 'Columns', self.formats['header'])
        worksheet.write(row, 5, 'Missing Values', self.formats['header'])
        row += 1
        
        for dataset_name, dataset in processed_data.items():
            worksheet.write(row, 0, dataset_name.title(), self.formats['cell_border'])
            worksheet.write(row, 1, len(dataset), self.formats['number'])
            worksheet.write(row, 2, dataset.index.min().strftime('%Y-%m-%d'), self.formats['date'])
            worksheet.write(row, 3, dataset.index.max().strftime('%Y-%m-%d'), self.formats['date'])
            worksheet.write(row, 4, len(dataset.columns), self.formats['number'])
            worksheet.write(row, 5, dataset.isnull().sum().sum(), self.formats['number'])
            row += 1
        
        # Column formatting
        worksheet.set_column('A:A', 15)
        worksheet.set_column('B:F', 12)
    
    def _create_exposure_analysis(self, writer, exposure_results):
        """Create exposure analysis worksheet"""
        worksheet = writer.book.add_worksheet('Exposure Analysis')
        
        # Title
        worksheet.merge_range('A1:D1', 'FBX Exposure Analysis', self.formats['title'])
        
        row = 2
        
        # FBX-Revenue relationship
        worksheet.write(row, 0, 'FBX-Revenue Relationship', self.formats['header'])
        row += 1
        
        fbx_revenue = exposure_results.get('fbx_revenue_analysis', {})
        
        metrics = [
            ('Correlation', fbx_revenue.get('correlation', 0), 'percentage'),
            ('Revenue Sensitivity (Beta)', fbx_revenue.get('linear_regression', {}).get('revenue_sensitivity', 0), 'percentage'),
            ('R-Squared', fbx_revenue.get('linear_regression', {}).get('r_squared', 0), 'percentage'),
            ('Statistical Significance (p-value)', fbx_revenue.get('statistical_tests', {}).get('beta_pvalue', 0), 'number')
        ]
        
        for metric_name, value, format_type in metrics:
            worksheet.write(row, 0, metric_name, self.formats['subheader'])
            if format_type == 'percentage':
                worksheet.write(row, 1, value, self.formats['percentage'])
            elif format_type == 'currency':
                worksheet.write(row, 1, value, self.formats['currency'])
            else:
                worksheet.write(row, 1, value, self.formats['number'])
            row += 1
        
        row += 1
        
        # Total exposure metrics
        worksheet.write(row, 0, 'Total Exposure Metrics', self.formats['header'])
        row += 1
        
        total_exposure = exposure_results.get('total_exposure', {})
        
        exposure_metrics = [
            ('Dollar Exposure', total_exposure.get('dollar_exposure', 0), 'currency'),
            ('FBX Volatility', total_exposure.get('fbx_volatility', 0), 'percentage'),
            ('Revenue VaR (95%)', total_exposure.get('parametric_var', 0), 'currency'),
            ('Expected Shortfall', total_exposure.get('expected_shortfall', 0), 'currency'),
            ('Annual Revenue at Risk', total_exposure.get('annual_revenue_at_risk', 0), 'currency')
        ]
        
        for metric_name, value, format_type in exposure_metrics:
            worksheet.write(row, 0, metric_name, self.formats['subheader'])
            if format_type == 'percentage':
                worksheet.write(row, 1, value, self.formats['percentage'])
            elif format_type == 'currency':
                worksheet.write(row, 1, value, self.formats['currency'])
            else:
                worksheet.write(row, 1, value, self.formats['number'])
            row += 1
        
        # Format columns
        worksheet.set_column('A:A', 25)
        worksheet.set_column('B:D', 15)
    
    def _create_hedge_ratios_analysis(self, writer, hedge_ratios):
        """Create hedge ratios analysis worksheet"""
        worksheet = writer.book.add_worksheet('Hedge Ratios')
        
        # Title
        worksheet.merge_range('A1:G1', 'Hedge Ratios Analysis', self.formats['title'])
        
        row = 2
        
        # Headers
        headers = ['Instrument', 'Optimal Ratio', 'Method', 'Effectiveness', 'Correlation', 'Volatility', 'Confidence']
        for col, header in enumerate(headers):
            worksheet.write(row, col, header, self.formats['header'])
        row += 1
        
        # Optimal hedge ratios
        optimal_ratios = hedge_ratios.get('optimal', {})
        
        for instrument, ratio_info in optimal_ratios.items():
            if isinstance(ratio_info, dict):
                col = 0
                worksheet.write(row, col, instrument, self.formats['cell_border'])
                col += 1
                
                optimal_ratio = ratio_info.get('optimal_hedge_ratio', 0)
                worksheet.write(row, col, optimal_ratio, self.formats['number'])
                col += 1
                
                best_method = ratio_info.get('best_method', 'N/A')
                worksheet.write(row, col, best_method, self.formats['cell_border'])
                col += 1
                
                # Get method details
                method_details = ratio_info.get('method_details', {})
                effectiveness = method_details.get('hedge_effectiveness', method_details.get('r_squared', 0))
                worksheet.write(row, col, effectiveness, self.formats['percentage'])
                col += 1
                
                correlation = method_details.get('correlation', 0)
                worksheet.write(row, col, correlation, self.formats['percentage'])
                col += 1
                
                volatility = method_details.get('instrument_volatility', method_details.get('volatility', 0))
                worksheet.write(row, col, volatility, self.formats['percentage'])
                col += 1
                
                confidence = ratio_info.get('confidence_score', 0)
                worksheet.write(row, col, confidence, self.formats['percentage'])
                
                row += 1
        
        # Format columns
        worksheet.set_column('A:A', 15)
        worksheet.set_column('B:G', 12)
    
    def _create_backtesting_results(self, writer, backtest_results):
        """Create backtesting results worksheet"""
        worksheet = writer.book.add_worksheet('Backtesting Results')
        
        # Title
        worksheet.merge_range('A1:J1', 'Backtesting Results', self.formats['title'])
        
        row = 2
        
        # Headers
        headers = ['Strategy', 'Total Return', 'Volatility', 'Sharpe Ratio', 'Max Drawdown', 
                  'VaR (95%)', 'Win Rate', 'Hedge Effectiveness', 'Trades', 'Final Value']
        for col, header in enumerate(headers):
            worksheet.write(row, col, header, self.formats['header'])
        row += 1
        
        # Results for each strategy
        for strategy_name, strategy_results in backtest_results.items():
            if not isinstance(strategy_results, dict) or strategy_name in ['comparison', 'hedge_effectiveness', 'risk_analysis', 'scenario_analysis']:
                continue
            
            col = 0
            worksheet.write(row, col, strategy_name.replace('_', ' ').title(), self.formats['cell_border'])
            col += 1
            
            # Total return
            total_return = strategy_results.get('total_return', 0)
            worksheet.write(row, col, total_return, self.formats['percentage'])
            col += 1
            
            # Volatility
            volatility = strategy_results.get('metrics', {}).get('volatility', 0)
            worksheet.write(row, col, volatility, self.formats['percentage'])
            col += 1
            
            # Sharpe ratio
            sharpe = strategy_results.get('metrics', {}).get('sharpe_ratio', 0)
            worksheet.write(row, col, sharpe, self.formats['number'])
            col += 1
            
            # Max drawdown
            max_dd = strategy_results.get('metrics', {}).get('max_drawdown', 0)
            worksheet.write(row, col, max_dd, self.formats['percentage'])
            col += 1
            
            # VaR
            var_95 = strategy_results.get('metrics', {}).get('var_historical', 0)
            worksheet.write(row, col, var_95, self.formats['percentage'])
            col += 1
            
            # Win rate
            win_rate = strategy_results.get('metrics', {}).get('win_rate', 0)
            worksheet.write(row, col, win_rate, self.formats['percentage'])
            col += 1
            
            # Hedge effectiveness
            hedge_eff = backtest_results.get('hedge_effectiveness', {}).get(strategy_name, {}).get('hedge_effectiveness', 0)
            worksheet.write(row, col, hedge_eff, self.formats['percentage'])
            col += 1
            
            # Number of trades
            trades = strategy_results.get('portfolio_metrics', {}).get('total_trades', 0)
            worksheet.write(row, col, trades, self.formats['number'])
            col += 1
            
            # Final value
            final_value = strategy_results.get('final_value', 0)
            worksheet.write(row, col, final_value, self.formats['currency'])
            
            row += 1
        
        # Format columns
        worksheet.set_column('A:A', 18)
        worksheet.set_column('B:J', 12)
    
    def _create_risk_analysis(self, writer, backtest_results):
        """Create risk analysis worksheet"""
        worksheet = writer.book.add_worksheet('Risk Analysis')
        
        # Title
        worksheet.merge_range('A1:F1', 'Risk Analysis', self.formats['title'])
        
        row = 2
        
        # Risk metrics comparison
        risk_analysis = backtest_results.get('risk_analysis', {})
        
        if risk_analysis:
            # Headers
            worksheet.write(row, 0, 'Strategy', self.formats['header'])
            worksheet.write(row, 1, 'VaR (95%)', self.formats['header'])
            worksheet.write(row, 2, 'Expected Shortfall', self.formats['header'])
            worksheet.write(row, 3, 'Tail Ratio', self.formats['header'])
            worksheet.write(row, 4, 'Extreme Negative', self.formats['header'])
            worksheet.write(row, 5, 'Extreme Positive', self.formats['header'])
            row += 1
            
            for strategy_name, strategy_risk in risk_analysis.items():
                if isinstance(strategy_risk, dict) and 'tail_risk' in strategy_risk:
                    tail_risk = strategy_risk['tail_risk']
                    
                    worksheet.write(row, 0, strategy_name.replace('_', ' ').title(), self.formats['cell_border'])
                    worksheet.write(row, 1, strategy_risk.get('var_metrics', {}).get('var_95', 0), self.formats['percentage'])
                    worksheet.write(row, 2, strategy_risk.get('var_metrics', {}).get('expected_shortfall', 0), self.formats['percentage'])
                    worksheet.write(row, 3, tail_risk.get('tail_ratio', 0), self.formats['number'])
                    worksheet.write(row, 4, tail_risk.get('extreme_negative', 0), self.formats['percentage'])
                    worksheet.write(row, 5, tail_risk.get('extreme_positive', 0), self.formats['percentage'])
                    row += 1
        
        # Format columns
        worksheet.set_column('A:A', 18)
        worksheet.set_column('B:F', 15)
    
    def _create_performance_comparison(self, writer, backtest_results):
        """Create performance comparison worksheet"""
        worksheet = writer.book.add_worksheet('Performance Comparison')
        
        # Title
        worksheet.merge_range('A1:E1', 'Performance Comparison vs Unhedged', self.formats['title'])
        
        row = 2
        
        comparison = backtest_results.get('comparison', {})
        
        if comparison:
            # Headers
            worksheet.write(row, 0, 'Strategy', self.formats['header'])
            worksheet.write(row, 1, 'Excess Return', self.formats['header'])
            worksheet.write(row, 2, 'Volatility Reduction', self.formats['header'])
            worksheet.write(row, 3, 'Information Ratio', self.formats['header'])
            worksheet.write(row, 4, 'Risk-Adjusted Return', self.formats['header'])
            row += 1
            
            for strategy_name, comparison_metrics in comparison.items():
                worksheet.write(row, 0, strategy_name.replace('_', ' ').title(), self.formats['cell_border'])
                worksheet.write(row, 1, comparison_metrics.get('excess_return', 0), self.formats['percentage'])
                worksheet.write(row, 2, comparison_metrics.get('volatility_reduction', 0), self.formats['percentage'])
                worksheet.write(row, 3, comparison_metrics.get('information_ratio', 0), self.formats['number'])
                worksheet.write(row, 4, comparison_metrics.get('risk_adjusted_return', 0), self.formats['number'])
                row += 1
        
        # Format columns
        worksheet.set_column('A:A', 18)
        worksheet.set_column('B:E', 15)
    
    def _create_scenario_analysis(self, writer, backtest_results):
        """Create scenario analysis worksheet"""
        worksheet = writer.book.add_worksheet('Scenario Analysis')
        
        # Title
        worksheet.merge_range('A1:F1', 'Scenario Analysis', self.formats['title'])
        
        row = 2
        
        scenario_analysis = backtest_results.get('scenario_analysis', {})
        
        if scenario_analysis:
            # Headers
            worksheet.write(row, 0, 'Strategy', self.formats['header'])
            worksheet.write(row, 1, 'Scenario', self.formats['header'])
            worksheet.write(row, 2, 'Avg Return', self.formats['header'])
            worksheet.write(row, 3, 'Volatility', self.formats['header'])
            worksheet.write(row, 4, 'Sharpe Ratio', self.formats['header'])
            worksheet.write(row, 5, 'Max Drawdown', self.formats['header'])
            row += 1
            
            for strategy_name, strategy_scenarios in scenario_analysis.items():
                if isinstance(strategy_scenarios, dict):
                    for scenario_name, scenario_metrics in strategy_scenarios.items():
                        if isinstance(scenario_metrics, dict):
                            worksheet.write(row, 0, strategy_name.replace('_', ' ').title(), self.formats['cell_border'])
                            worksheet.write(row, 1, scenario_name.replace('_', ' ').title(), self.formats['cell_border'])
                            worksheet.write(row, 2, scenario_metrics.get('avg_return', 0), self.formats['percentage'])
                            worksheet.write(row, 3, scenario_metrics.get('volatility', 0), self.formats['percentage'])
                            worksheet.write(row, 4, scenario_metrics.get('sharpe_ratio', 0), self.formats['number'])
                            worksheet.write(row, 5, scenario_metrics.get('max_drawdown', 0), self.formats['percentage'])
                            row += 1
        
        # Format columns
        worksheet.set_column('A:B', 18)
        worksheet.set_column('C:F', 12)
    
    def _create_detailed_data(self, writer, processed_data):
        """Create detailed data worksheets"""
        
        # FBX data
        if 'fbx' in processed_data:
            fbx_data = processed_data['fbx'].copy()
            fbx_data.to_excel(writer, sheet_name='FBX Data', index=True)
        
        # Instruments data (sample)
        if 'instruments' in processed_data:
            instruments_data = processed_data['instruments'].copy()
            # Limit to prevent Excel from becoming too large
            if len(instruments_data) > 1000:
                instruments_data = instruments_data.tail(1000)
            instruments_data.to_excel(writer, sheet_name='Instruments Data', index=True)
        
        # Revenue data
        if 'revenue' in processed_data:
            revenue_data = processed_data['revenue'].copy()
            revenue_data.to_excel(writer, sheet_name='Revenue Data', index=True)
    
    def _add_charts_to_report(self, writer, processed_data, backtest_results, ml_results=None):
        """Add charts to the report"""
        try:
            # Performance chart
            if backtest_results:
                perf_chart = self.chart_generator.create_performance_chart(backtest_results)
                # Note: xlsxwriter doesn't support plotly charts directly
                # In a production environment, you'd save charts as images and insert them
                
        except Exception as e:
            self.logger.error(f"Error adding charts to report: {e}")
    
    def _get_best_strategy(self, backtest_results):
        """Get the best performing strategy"""
        best_strategy = None
        best_score = -np.inf
        
        for strategy_name, strategy_results in backtest_results.items():
            if not isinstance(strategy_results, dict) or strategy_name in ['comparison', 'hedge_effectiveness', 'risk_analysis', 'scenario_analysis']:
                continue
            
            # Calculate composite score
            total_return = strategy_results.get('total_return', 0)
            volatility = strategy_results.get('metrics', {}).get('volatility', 1)
            max_drawdown = strategy_results.get('metrics', {}).get('max_drawdown', 0)
            
            # Risk-adjusted return score
            score = total_return / volatility - abs(max_drawdown)
            
            if score > best_score:
                best_score = score
                best_strategy = (strategy_name, strategy_results)
        
        return best_strategy
    
    def _generate_recommendations(self, exposure_results, hedge_ratios, backtest_results):
        """Generate recommendations based on analysis"""
        recommendations = []
        
        # Exposure-based recommendations
        total_exposure = exposure_results.get('total_exposure', {})
        revenue_sensitivity = total_exposure.get('revenue_sensitivity', 0)
        
        if abs(revenue_sensitivity) > 0.5:
            recommendations.append("High revenue sensitivity to FBX detected. Consider implementing hedging strategy.")
        
        # Hedge effectiveness recommendations
        hedge_effectiveness = backtest_results.get('hedge_effectiveness', {})
        best_hedge = max(hedge_effectiveness.items(), 
                        key=lambda x: x[1].get('hedge_effectiveness', 0) if isinstance(x[1], dict) else 0,
                        default=(None, None))
        
        if best_hedge[0] and best_hedge[1]:
            effectiveness = best_hedge[1].get('hedge_effectiveness', 0)
            if effectiveness > 0.3:
                recommendations.append(f"Recommend {best_hedge[0].replace('_', ' ')} strategy with {effectiveness:.1%} hedge effectiveness.")
        
        # Risk recommendations
        unhedged_vol = backtest_results.get('unhedged', {}).get('metrics', {}).get('volatility', 0)
        if unhedged_vol > 0.2:
            recommendations.append("High volatility detected. Hedging could significantly reduce risk.")
        
        # Position sizing recommendations
        optimal_ratios = hedge_ratios.get('optimal', {})
        high_confidence_instruments = [inst for inst, info in optimal_ratios.items() 
                                     if isinstance(info, dict) and info.get('confidence_score', 0) > 0.7]
        
        if high_confidence_instruments:
            recommendations.append(f"Focus on high-confidence instruments: {', '.join(high_confidence_instruments[:3])}")
        
        # Diversification recommendations
        if len(optimal_ratios) > 1:
            recommendations.append("Consider diversifying hedge across multiple instruments to reduce concentration risk.")
        
        return recommendations
    
    def _create_ml_analysis(self, writer, ml_results: Dict[str, Any]):
        """Create ML analysis worksheet"""
        worksheet = writer.book.add_worksheet('ML Analysis')
        
        # Title
        worksheet.merge_range('A1:E1', 'Machine Learning Analysis', self.formats['title'])
        
        row = 3
        
        # ML Pipeline Summary
        worksheet.write(row, 0, 'ML Pipeline Summary', self.formats['header'])
        row += 2
        
        summary_data = [
            ['Status', ml_results.get('status', 'N/A')],
            ['Total Time (seconds)', ml_results.get('total_time', 0)],
            ['Features Created', ml_results.get('features_created', 0)],
            ['Features Selected', ml_results.get('features_selected', 0)],
            ['Models Trained', ml_results.get('models_trained', 0)],
            ['Ensembles Created', ml_results.get('ensembles_created', 0)],
            ['Best Single Model', ml_results.get('best_single_model', 'N/A')],
            ['Best Ensemble', ml_results.get('best_ensemble', 'N/A')]
        ]
        
        for data_row in summary_data:
            worksheet.write(row, 0, data_row[0], self.formats['subheader'])
            worksheet.write(row, 1, data_row[1], self.formats['number'])
            row += 1
        
        row += 2
        
        # Model Performance Summary
        if 'model_performances' in ml_results:
            worksheet.write(row, 0, 'Model Performance Summary', self.formats['header'])
            row += 2
            
            # Headers
            headers = ['Model', 'R² Score', 'MSE', 'MAE']
            for col, header in enumerate(headers):
                worksheet.write(row, col, header, self.formats['subheader'])
            row += 1
            
            # Data
            for model_name, metrics in ml_results['model_performances'].items():
                worksheet.write(row, 0, model_name, self.formats['text'])
                worksheet.write(row, 1, metrics.get('r2_score', 0), self.formats['number'])
                worksheet.write(row, 2, metrics.get('mse', 0), self.formats['number'])
                worksheet.write(row, 3, metrics.get('mae', 0), self.formats['number'])
                row += 1
        
        # Set column widths
        worksheet.set_column('A:A', 20)
        worksheet.set_column('B:E', 15)
    
    def _create_ml_predictions(self, writer, ml_results: Dict[str, Any]):
        """Create ML predictions worksheet"""
        worksheet = writer.book.add_worksheet('ML Predictions')
        
        # Title
        worksheet.merge_range('A1:G1', 'Machine Learning Predictions', self.formats['title'])
        
        row = 3
        
        # Test Set Predictions
        if 'predictions' in ml_results and 'test_predictions' in ml_results['predictions']:
            worksheet.write(row, 0, 'Test Set Predictions', self.formats['header'])
            row += 2
            
            test_predictions = ml_results['predictions']['test_predictions']
            
            # Get first model's dates for reference
            first_model_name = list(test_predictions.keys())[0]
            dates = test_predictions[first_model_name].get('dates', [])
            
            # Create headers
            headers = ['Date'] + list(test_predictions.keys())
            for col, header in enumerate(headers):
                worksheet.write(row, col, header, self.formats['subheader'])
            row += 1
            
            # Write prediction data
            for i, date in enumerate(dates[:100]):  # Limit to first 100 predictions
                worksheet.write(row, 0, date, self.formats['text'])
                
                for col, model_name in enumerate(list(test_predictions.keys()), 1):
                    predictions = test_predictions[model_name].get('predictions', [])
                    if i < len(predictions):
                        worksheet.write(row, col, predictions[i], self.formats['number'])
                
                row += 1
        
        row += 2
        
        # Future Predictions
        if 'predictions' in ml_results and 'future_predictions' in ml_results['predictions']:
            worksheet.write(row, 0, 'Future Predictions', self.formats['header'])
            row += 2
            
            future_predictions = ml_results['predictions']['future_predictions']
            
            # Headers
            headers = ['Horizon'] + list(list(future_predictions.values())[0].keys())
            for col, header in enumerate(headers):
                worksheet.write(row, col, header, self.formats['subheader'])
            row += 1
            
            # Data
            for horizon, predictions in future_predictions.items():
                worksheet.write(row, 0, horizon, self.formats['text'])
                
                for col, model_name in enumerate(list(predictions.keys()), 1):
                    worksheet.write(row, col, predictions[model_name], self.formats['number'])
                
                row += 1
        
        # Set column widths
        worksheet.set_column('A:A', 20)
        worksheet.set_column('B:G', 15)
    
    def _create_ml_model_comparison(self, writer, ml_results: Dict[str, Any]):
        """Create ML model comparison worksheet"""
        worksheet = writer.book.add_worksheet('ML Model Comparison')
        
        # Title
        worksheet.merge_range('A1:F1', 'Machine Learning Model Comparison', self.formats['title'])
        
        row = 3
        
        # Model Comparison Table
        if 'model_performances' in ml_results:
            worksheet.write(row, 0, 'Model Performance Comparison', self.formats['header'])
            row += 2
            
            # Headers
            headers = ['Model', 'R² Score', 'MSE', 'RMSE', 'MAE', 'Rank']
            for col, header in enumerate(headers):
                worksheet.write(row, col, header, self.formats['subheader'])
            row += 1
            
            # Sort models by R² score
            performances = ml_results['model_performances']
            sorted_models = sorted(performances.items(), 
                                 key=lambda x: x[1].get('r2_score', -np.inf), 
                                 reverse=True)
            
            # Data
            for rank, (model_name, metrics) in enumerate(sorted_models, 1):
                worksheet.write(row, 0, model_name, self.formats['text'])
                worksheet.write(row, 1, metrics.get('r2_score', 0), self.formats['number'])
                worksheet.write(row, 2, metrics.get('mse', 0), self.formats['number'])
                worksheet.write(row, 3, np.sqrt(metrics.get('mse', 0)), self.formats['number'])
                worksheet.write(row, 4, metrics.get('mae', 0), self.formats['number'])
                worksheet.write(row, 5, rank, self.formats['number'])
                row += 1
        
        row += 2
        
        # Model Recommendations
        worksheet.write(row, 0, 'Model Recommendations', self.formats['header'])
        row += 2
        
        recommendations = self._generate_ml_recommendations(ml_results)
        for recommendation in recommendations:
            worksheet.write(row, 0, f"• {recommendation}", self.formats['text'])
            row += 1
        
        # Set column widths
        worksheet.set_column('A:A', 25)
        worksheet.set_column('B:F', 15)
    
    def _generate_ml_recommendations(self, ml_results: Dict[str, Any]) -> List[str]:
        """Generate ML-based recommendations"""
        recommendations = []
        
        # Best model recommendation
        best_model = ml_results.get('best_single_model')
        if best_model:
            recommendations.append(f"Use {best_model} as primary prediction model")
        
        # Ensemble recommendation
        best_ensemble = ml_results.get('best_ensemble')
        if best_ensemble:
            recommendations.append(f"Consider {best_ensemble} ensemble for improved accuracy")
        
        # Performance-based recommendations
        if 'model_performances' in ml_results:
            performances = ml_results['model_performances']
            
            # Find models with R² > 0.5
            good_models = [name for name, metrics in performances.items() 
                          if metrics.get('r2_score', 0) > 0.5]
            
            if good_models:
                recommendations.append(f"Models with good performance (R² > 0.5): {', '.join(good_models[:3])}")
            else:
                recommendations.append("Consider feature engineering improvements - low R² scores observed")
        
        # Feature importance recommendation
        if ml_results.get('features_selected', 0) > 0:
            recommendations.append("Review feature importance analysis for key predictors")
        
        return recommendations
    
    def _update_executive_summary_with_ml(self, writer, ml_results: Dict[str, Any]):
        """Update executive summary to include ML results"""
        # This would be called from _create_executive_summary if ML results exist
        pass
